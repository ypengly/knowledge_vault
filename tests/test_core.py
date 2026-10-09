import json
import unittest
from pathlib import Path

from knowledge_vault import ingest, qa
from knowledge_vault.chunker import chunk_text
from knowledge_vault.extract import ExtractionError, extract, html_to_text, split_sections
from knowledge_vault.llm import Provider
from knowledge_vault.search import Filters
from knowledge_vault.textutil import coverage, highlight_html

from .helpers import make_vault, write

RETENTION = ("Customer retention is the strongest driver of profit. Our research found that "
             "a five percent increase in customer retention raised profits by twenty five percent. "
             "Loyalty programs and onboarding emails reduced churn among new customers.")
OTHER = "The garden needs watering twice a week. Tomatoes prefer full sun and well drained soil."


class FakeLLM(Provider):
    name, model, base_url = "Fake", "fake-1", "http://localhost:1"

    def __init__(self, reply, external=False):
        self.reply, self._ext, self.calls = reply, external, 0

    @property
    def is_external(self):
        return self._ext

    def generate(self, system, user):
        self.calls += 1
        self.last_user = user
        return self.reply


class Base(unittest.TestCase):
    def setUp(self):
        self.v, self.tmp = make_vault()
        self.a = ingest.import_file(self.v, write(self.tmp, "Business Research.txt", RETENTION))
        self.b = ingest.import_file(self.v, write(self.tmp, "garden.md", OTHER))


class TestUnits(unittest.TestCase):
    def test_chunker_overlap_and_bounds(self):
        text = " ".join(f"w{i}" for i in range(100))
        chunks = chunk_text(text, 40, 10)
        self.assertEqual(chunks[0].split()[0], "w0")
        self.assertEqual(chunks[1].split()[0], "w30")
        self.assertEqual(chunks[-1].split()[-1], "w99")
        self.assertEqual(chunk_text("short text", 40, 10), ["short text"])
        self.assertEqual(chunk_text("   ", 40, 10), [])

    def test_sections_and_csv_labels(self):
        units = split_sections("a\n\n" + "b" * 50 + "\n\nc", target=40)
        self.assertGreater(len(units), 1)
        self.assertEqual(units[0].label, "Section 1")

    def test_csv_extraction(self):
        v, tmp = make_vault()
        p = write(tmp, "d.csv", "name,city\n" + "\n".join(f"p{i},c{i}" for i in range(100)))
        res = extract(p, "csv")
        self.assertEqual(res.units[0].label, "Rows 1\u201340")
        self.assertIn("city: c3", res.units[0].text)

    def test_empty_and_unsupported(self):
        v, tmp = make_vault()
        with self.assertRaises(ExtractionError):
            extract(write(tmp, "e.txt", "   "), "txt")
        with self.assertRaises(ExtractionError):
            ingest.register_file(v, write(tmp, "x.exe", "zz"))

    def test_docx_roundtrip(self):
        docx = __import__("pytest").importorskip if False else None
        try:
            import docx as d
        except ImportError:
            self.skipTest("python-docx missing")
        v, tmp = make_vault()
        doc = d.Document(); doc.add_paragraph("Quarterly onboarding improved retention.")
        t = doc.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Metric"; t.rows[0].cells[1].text = "Churn"
        doc.save(tmp / "r.docx")
        res = extract(tmp / "r.docx", "docx")
        self.assertIn("retention", res.units[0].text)
        self.assertIn("Metric | Churn", res.units[0].text)

    def test_pdf_pages(self):
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF missing")
        v, tmp = make_vault()
        doc = fitz.open()
        for i in range(3):
            doc.new_page().insert_text((72, 72), f"Page {i + 1} talks about retention strategy number {i}")
        doc.save(tmp / "r.pdf")
        did = ingest.import_file(v, tmp / "r.pdf")
        self.assertEqual(v.db.get_document(did)["pages"], 3)
        hits = v.search("retention strategy number 2", "keyword")
        self.assertEqual(hits[0].label, "Page 3")

    def test_html_to_text_skips_scripts(self):
        title, text = html_to_text("<html><head><title>T</title></head><body><script>bad()</script>"
                                   "<p>Hello</p><p>World</p></body></html>")
        self.assertEqual(title, "T")
        self.assertNotIn("bad", text)
        self.assertIn("Hello", text)

    def test_highlight_escapes_and_marks(self):
        out = highlight_html("Retention <b>rates</b> improve", "customer retention rate")
        self.assertIn("<mark>Retention</mark>", out)
        self.assertIn("&lt;b&gt;", out)
        self.assertNotIn("<b>", out)
        self.assertGreater(coverage("customer retention", RETENTION), 0.9)


class TestPipeline(Base):
    def test_status_and_metadata(self):
        d = self.v.db.get_document(self.a)
        self.assertEqual(d["status"], "ready")
        self.assertEqual(d["file_type"], "txt")
        self.assertGreater(d["size"], 0)

    def test_duplicate_import_is_skipped(self):
        again = ingest.import_file(self.v, self.tmp / "Business Research.txt")
        self.assertEqual(again, self.a)
        self.assertEqual(self.v.db.stats(self.v.files_dir)["documents"], 2)

    def test_failed_document_is_reported(self):
        p = write(self.tmp, "blank.txt", "  \n ")
        did = ingest.import_file(self.v, p)
        d = self.v.db.get_document(did)
        self.assertEqual(d["status"], "failed")
        self.assertIn("No extractable text", d["error"])

    def test_keyword_semantic_hybrid(self):
        for mode in ("keyword", "semantic", "hybrid"):
            hits = self.v.search("customer retention", mode)
            self.assertTrue(hits, mode)
            self.assertEqual(hits[0].filename, "Business Research.txt", mode)
        self.assertEqual(self.v.search("zzzxqv", "keyword"), [])

    def test_filters(self):
        coll = self.v.db.q("SELECT id FROM collections WHERE name='Research'")[0]["id"]
        self.v.db.assign(self.a, coll)
        q = "customer retention tomatoes"
        self.assertEqual({h.doc_id for h in self.v.search(q, "keyword", Filters(collection_id=coll))}, {self.a})
        self.assertEqual({h.doc_id for h in self.v.search(q, "semantic", Filters(collection_id=coll))}, {self.a})
        self.assertEqual({h.doc_id for h in self.v.search(q, "keyword", Filters(file_types=["md"]))}, {self.b})
        self.assertEqual(self.v.search(q, "keyword", Filters(date_from="2999-01-01")), [])
        self.assertTrue(self.v.search(q, "keyword", Filters(date_to="2999-01-01")))

    def test_delete_document_removes_everything(self):
        self.v.delete_document(self.a)
        self.assertEqual(self.v.search("customer retention", "keyword"), [])
        self.assertEqual(self.v.search("customer retention", "semantic"), [])
        self.assertFalse(any(self.tmp.joinpath("files").glob("*Business*")))

    def test_index_survives_restart(self):
        from knowledge_vault.core import Vault
        v2 = Vault(self.tmp, self.v.settings)
        self.assertEqual(len(v2.index), len(self.v.index))
        self.assertTrue(v2.search("tomatoes sun", "semantic"))

    def test_collections_crud(self):
        names = [c["name"] for c in self.v.db.collections()]
        for n in ("Work", "School", "Research", "Business", "Personal"):
            self.assertIn(n, names)
        cid = self.v.db.add_collection("Clients")
        self.v.db.rename_collection(cid, "Customers")
        self.assertIn("Customers", [c["name"] for c in self.v.db.collections()])
        self.v.db.delete_collection(cid)
        self.assertNotIn("Customers", [c["name"] for c in self.v.db.collections()])
        work = [c for c in self.v.db.collections() if c["name"] == "Work"][0]["id"]
        self.v.db.delete_collection(work)  # built-ins are protected
        self.assertIn("Work", [c["name"] for c in self.v.db.collections()])

    def test_notes(self):
        nid = self.v.db.save_note(None, "Retention idea", "# Plan\nCall churned customers",
                                  "Research", ["#CRM", "ideas"], [self.a])
        n = self.v.db.get_note(nid)
        self.assertEqual(n["tags"], ["crm", "ideas"])
        self.assertEqual(n["doc_ids"], [self.a])
        self.assertEqual(len(self.v.db.list_notes(text="churned")), 1)
        self.assertEqual(len(self.v.db.list_notes(tag="crm", folder="Research")), 1)
        self.assertEqual(self.v.db.notes_for_document(self.a)[0]["title"], "Retention idea")
        self.v.db.save_note(nid, "Renamed", "x", "", [], [])
        self.assertEqual(self.v.db.get_note(nid)["title"], "Renamed")
        self.v.db.delete_note(nid)
        self.assertIsNone(self.v.db.get_note(nid))


class TestAsk(Base):
    def test_insufficient_when_nothing_relevant(self):
        ans = qa.ask(self.v, "What is the capital of Mongolia?", FakeLLM("Ulaanbaatar [1]"))
        self.assertEqual(ans.text, qa.INSUFFICIENT)
        self.assertEqual(ans.citations, [])
        self.assertFalse(ans.grounded)

    def test_llm_not_called_without_evidence(self):
        llm = FakeLLM("made up")
        qa.ask(self.v, "quantum chromodynamics lattice", llm)
        self.assertEqual(llm.calls, 0)

    def test_extractive_mode_has_citations_and_no_fabrication(self):
        ans = qa.ask(self.v, "What did my research say about customer retention?")
        self.assertEqual(ans.mode, "extractive")
        self.assertTrue(ans.grounded)
        self.assertEqual(ans.citations[0].filename, "Business Research.txt")
        self.assertIn("nothing was generated", ans.text)

    def test_llm_answer_citations_validated(self):
        llm = FakeLLM("Retention raised profit by 25% [1]. Invented [9].")
        ans = qa.ask(self.v, "customer retention profit", llm)
        self.assertEqual([c.n for c in ans.citations], [1])
        self.assertTrue(ans.verified)
        self.assertIn("Business Research.txt", llm.last_user)
        self.assertNotIn("tomatoes", llm.last_user)

    def test_uncited_llm_answer_is_flagged(self):
        ans = qa.ask(self.v, "customer retention profit", FakeLLM("Profit goes up."))
        self.assertFalse(ans.verified)
        self.assertTrue(ans.citations)
        self.assertIn("unverified", ans.note)

    def test_llm_can_decline(self):
        ans = qa.ask(self.v, "customer retention profit", FakeLLM(qa.INSUFFICIENT))
        self.assertEqual(ans.text, qa.INSUFFICIENT)
        self.assertEqual(ans.citations, [])

    def test_external_requires_confirmation(self):
        llm = FakeLLM("ok [1]", external=True)
        with self.assertRaises(qa.ExternalNotConfirmed):
            qa.ask(self.v, "customer retention profit", llm)
        self.assertEqual(llm.calls, 0)
        qa.ask(self.v, "customer retention profit", llm, allow_external=True)
        self.assertEqual(llm.calls, 1)

    def test_history_logged(self):
        self.v.search("retention")
        qa.ask(self.v, "customer retention")
        self.assertEqual(self.v.db.recent_searches()[0]["query"], "retention")
        row = self.v.db.q("SELECT citations FROM questions")[0]
        self.assertTrue(json.loads(row["citations"]))

    def test_privacy_status_local_by_default(self):
        text, external = self.v.privacy_status()
        self.assertFalse(external)
        self.assertIn("local", text.lower())

    def test_no_hardcoded_keys_in_source(self):
        import re
        for p in Path(__file__).parent.parent.joinpath("knowledge_vault").rglob("*.py"):
            self.assertIsNone(re.search(r"sk-[A-Za-z0-9]{20,}", p.read_text("utf-8")), p)


if __name__ == "__main__":
    unittest.main()
