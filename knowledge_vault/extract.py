"""Text extraction. Every extractor returns Units: (number, label, text) so that
citations can point at an exact page / section / row range."""
from __future__ import annotations

import csv
import html.parser
import io
import re
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlparse

SUPPORTED = {
    ".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "md", ".markdown": "md",
    ".csv": "csv", ".png": "image", ".jpg": "image", ".jpeg": "image",
    ".tif": "image", ".tiff": "image", ".bmp": "image", ".webp": "image",
}


class ExtractionError(Exception):
    pass


@dataclass
class Unit:
    num: int
    label: str
    text: str


@dataclass
class Extracted:
    units: list[Unit]
    page_count: int
    warnings: list[str] = field(default_factory=list)


def file_type_for(path: Path) -> str | None:
    return SUPPORTED.get(path.suffix.lower())


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def split_sections(text: str, target: int = 3000, label: str = "Section") -> list[Unit]:
    """Virtual 'pages' for formats that have none, split on paragraph boundaries."""
    paras = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
    units, buf = [], ""
    for p in paras:
        if buf and len(buf) + len(p) > target:
            units.append(buf.strip())
            buf = ""
        buf += p + "\n\n"
    if buf.strip():
        units.append(buf.strip())
    return [Unit(i + 1, f"{label} {i + 1}", t) for i, t in enumerate(units)]


# ---------------------------------------------------------------- OCR
def ocr_image(img, lang: str = "eng", tesseract_cmd: str = "") -> str:
    try:
        import pytesseract
    except ImportError as e:
        raise ExtractionError("OCR needs 'pytesseract' (pip install pytesseract) and the "
                              "Tesseract engine (see README).") from e
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    try:
        return pytesseract.image_to_string(img, lang=lang)
    except pytesseract.TesseractNotFoundError as e:
        raise ExtractionError("Tesseract OCR engine not found. Install it and set its path "
                              "in Settings (see README).") from e


# ---------------------------------------------------------------- formats
def extract_pdf(path: Path, ocr_lang="eng", tesseract_cmd="") -> Extracted:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ExtractionError("PyMuPDF is required for PDFs (pip install pymupdf).") from e
    units, warnings = [], []
    try:
        doc = fitz.open(str(path))
    except Exception as e:
        raise ExtractionError(f"Could not open PDF: {e}") from e
    with doc:
        if doc.needs_pass:
            raise ExtractionError("PDF is password-protected.")
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if len(text) < 20:  # probably scanned -> OCR fallback
                try:
                    from PIL import Image
                    pix = page.get_pixmap(dpi=200)
                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    text = ocr_image(img, ocr_lang, tesseract_cmd).strip() or text
                except ExtractionError as e:
                    if not any("OCR" in w or "Tesseract" in w for w in warnings):
                        warnings.append(f"Scanned page(s) skipped: {e}")
                except Exception:
                    pass
            units.append(Unit(i, f"Page {i}", text))
        count = doc.page_count
    return Extracted(units, count, warnings)


def extract_docx(path: Path, **_) -> Extracted:
    try:
        import docx
    except ImportError as e:
        raise ExtractionError("python-docx is required for DOCX (pip install python-docx).") from e
    try:
        d = docx.Document(str(path))
    except Exception as e:
        raise ExtractionError(f"Could not open DOCX: {e}") from e
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    units = split_sections("\n\n".join(parts))
    return Extracted(units, len(units), ["DOCX has no fixed pages; citations use sections."])


def extract_text(path: Path, **_) -> Extracted:
    units = split_sections(read_text_file(path))
    return Extracted(units, len(units))


def extract_csv(path: Path, rows_per_unit: int = 40, **_) -> Extracted:
    text = read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return Extracted([], 0)
    header, body = rows[0], rows[1:]
    units = []
    for start in range(0, max(len(body), 1), rows_per_unit):
        block = body[start: start + rows_per_unit]
        lines = ["; ".join(f"{h}: {v}" for h, v in zip(header, r) if v.strip()) for r in block]
        label = f"Rows {start + 1}\u2013{start + len(block)}" if block else "Header"
        units.append(Unit(len(units) + 1, label, "\n".join(l for l in lines if l) or "; ".join(header)))
    return Extracted(units, len(units))


def extract_image(path: Path, ocr_lang="eng", tesseract_cmd="", **_) -> Extracted:
    try:
        from PIL import Image
        img = Image.open(path)
    except Exception as e:
        raise ExtractionError(f"Could not open image: {e}") from e
    text = ocr_image(img, ocr_lang, tesseract_cmd).strip()
    if not text:
        raise ExtractionError("OCR found no text in this image.")
    return Extracted([Unit(1, "Image", text)], 1)


EXTRACTORS = {"pdf": extract_pdf, "docx": extract_docx, "txt": extract_text, "md": extract_text,
              "url": extract_text, "csv": extract_csv, "image": extract_image}


def extract(path: Path, file_type: str, ocr_lang="eng", tesseract_cmd="") -> Extracted:
    fn = EXTRACTORS.get(file_type)
    if not fn:
        raise ExtractionError(f"Unsupported file type: {file_type}")
    res = fn(path, ocr_lang=ocr_lang, tesseract_cmd=tesseract_cmd)
    if not any(u.text.strip() for u in res.units):
        raise ExtractionError("No extractable text found.")
    return res


# ---------------------------------------------------------------- URLs
class _TextHTML(html.parser.HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "head", "nav", "footer"}
    BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article"}

    def __init__(self):
        super().__init__()
        self.out, self.title, self._skip, self._in_title = [], "", 0, False

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True
        if tag in self.SKIP:
            self._skip += 1
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self._skip:
            self.out.append(data)


def html_to_text(raw_html: str) -> tuple[str, str]:
    p = _TextHTML()
    p.feed(raw_html)
    text = re.sub(r"[ \t]+", " ", "".join(p.out))
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    return p.title.strip(), text


def fetch_url(url: str, timeout: int = 20) -> tuple[str, str]:
    """Explicit, user-initiated download of ONE page. Returns (title, text)."""
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise ExtractionError("Please enter a full http(s) URL.")
    req = urllib.request.Request(url, headers={"User-Agent": "KnowledgeVault/1.0 (+local app)"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            ctype = r.headers.get("Content-Type", "")
            raw = r.read(8_000_000)
            charset = r.headers.get_content_charset() or "utf-8"
    except Exception as e:
        raise ExtractionError(f"Could not download page: {e}") from e
    if "html" not in ctype and "text" not in ctype:
        raise ExtractionError(f"URL is not a web page ({ctype or 'unknown type'}). Download it and import the file.")
    title, text = html_to_text(raw.decode(charset, errors="replace"))
    if not text:
        raise ExtractionError("No readable text found at that URL.")
    return title or parsed.netloc, text
